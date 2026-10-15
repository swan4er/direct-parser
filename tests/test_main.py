"""Тесты без сети: образцы в fixtures/ — ужатые настоящие страницы yandex.ru/search/ads
(14.10.2026): ads_page.html — первые два объявления по «купить духи»,
no_ads_page.html — ответ «ничего не нашлось»."""

import csv
import io
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

import main

ROOT = Path(__file__).parent.parent
FIXTURES = Path(__file__).parent / "fixtures"
ADS_PAGE = (FIXTURES / "ads_page.html").read_text(encoding="utf-8")
NO_ADS_PAGE = (FIXTURES / "no_ads_page.html").read_text(encoding="utf-8")
CAPTCHA_PAGE = '<form action="/checkcaptcha?key=x"></form>'


def ads_html(*domains):
    """Страница с объявлениями заданных доменов — в той же разметке, что ads_page.html."""
    items = "".join(
        f'<li class="serp-item serp-item_card" data-cid="{i}">'
        f'<a class="Link OrganicTitle-Link" href="https://yabs.yandex.ru/count/{domain}">'
        f'<h2 class="OrganicTitle-LinkText">Объявление {domain}</h2></a>'
        f'<div class="Path Organic-Path"><a class="Link Path-Item" href="#"><b>{domain}</b>'
        f'<span class="Path-Separator">›</span>страница</a></div>'
        f'<span class="OrganicTextContentSpan">Текст {domain}</span></li>'
        for i, domain in enumerate(domains))
    return f'<ul id="search-result">{items}</ul>'


class FakeResponse:
    def __init__(self, text, url="https://yandex.ru/search/ads"):
        self.text, self.url = text, url

    def raise_for_status(self):
        pass


class FakeSession:
    """Отдаёт страницы по номеру p; всё, что за пределами pages, — «ничего не нашлось»."""

    def __init__(self, pages, url="https://yandex.ru/search/ads"):
        self.pages, self.url, self.calls = pages, url, []

    def get(self, url, params, timeout):
        self.calls.append(params)
        page = params["p"]
        return FakeResponse(self.pages[page] if page < len(self.pages) else NO_ADS_PAGE, self.url)


class QueueSession(FakeSession):
    """Отдаёт страницы по очереди вызовов (удобно для прогонов); потом — «ничего не нашлось»."""

    def __init__(self, *texts):
        super().__init__([])
        self.texts = list(texts)

    def get(self, url, params, timeout):
        self.calls.append(params)
        return FakeResponse(self.texts.pop(0) if self.texts else NO_ADS_PAGE, self.url)


class FakeTTY(io.StringIO):
    def isatty(self):
        return True


@pytest.fixture(autouse=True)
def isolated_results(monkeypatch, tmp_path):
    """Результаты «по умолчанию» — во временную папку, а не в results/ проекта."""
    monkeypatch.setattr(main, "RESULTS_DIR", tmp_path.resolve() / "results")


@pytest.fixture
def no_sleep(monkeypatch):
    monkeypatch.setattr(main.time, "sleep", lambda seconds: None)


def use_session(monkeypatch, session):
    monkeypatch.setattr(main, "make_session", lambda proxy=None: session)
    return session


def run_cli(monkeypatch, session, *argv):
    use_session(monkeypatch, session)
    return main.main([*argv, "--delay", "0"])


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def answers(*values):
    """Подменяет input(): отдаёт ответы по очереди, а когда они кончились — EOF."""
    queue = iter(values)

    def ask(prompt=""):
        try:
            return next(queue)
        except StopIteration:
            raise EOFError from None
    return ask


# ---------- разбор страницы ----------

def test_parse_real_ads_page():
    result = main.parse_page(ADS_PAGE)
    assert result["status"] == "ok"
    assert [ad["domain"] for ad in result["ads"]] == ["randewoo.ru", "impexoils.ru"]
    first = result["ads"][0]
    assert first["title"] == "Парфюм в интернет-магазине Randewoo — скидки до 55%"
    assert first["description"].startswith("Раскройте свой потенциал")
    # неразрывные пробелы Яндекса заменены обычными
    assert first["sitelinks"] == ["Aroma Box", "Бестселлеры", "Конструктор Aroma Box", "SALE"]
    assert first["label"] == "Промо"
    assert first["click_url"].startswith("https://yabs.yandex.ru/count/")


def test_parse_real_no_ads_page():
    assert main.parse_page(NO_ADS_PAGE) == {"status": "no_ads", "ads": []}


def test_parse_captcha():
    assert main.parse_page("<html></html>", "https://yandex.ru/showcaptcha?x")["status"] == "captcha"
    assert main.parse_page(CAPTCHA_PAGE)["status"] == "captcha"


def test_parse_unknown_markup():
    assert main.parse_page('<ul id="search-result"></ul>')["status"] == "unexpected"


def test_parse_ad_with_missing_parts_does_not_crash():
    html = ('<ul id="search-result"><li class="serp-item" data-cid="0">'
            '<a class="Link" href="https://yabs.yandex.ru/count/x">'
            '<h2 class="OrganicTitle-LinkText">Только заголовок</h2></a></li></ul>')
    ad = main.parse_page(html)["ads"][0]
    assert ad == {"domain": "", "title": "Только заголовок", "description": "",
                  "sitelinks": [], "label": "", "click_url": "https://yabs.yandex.ru/count/x"}


def test_synthetic_page_parses_like_real_one():
    assert [ad["domain"] for ad in main.parse_page(ads_html("a.ru", "b.ru"))["ads"]] == [
        "a.ru", "b.ru"]


# ---------- регионы и запросы ----------

@pytest.mark.parametrize("value, code", [
    ("36", 36), ("Ставрополь", 36), ("  ростов на дону ", 39), ("Ростов-на-Дону", 39),
    ("Орел", 10), ("спб", 2), ("Москва", 213), ("москва и область", 1), ("99999", 99999),
])
def test_parse_region(value, code):
    assert main.parse_region(value) == code


@pytest.mark.parametrize("value", ["Атлантида", "0", ""])
def test_parse_region_rejects_unknown(value):
    with pytest.raises(ValueError, match="aistudio.yandex.ru"):
        main.parse_region(value)


def test_read_queries_from_args_and_file(tmp_path):
    path = tmp_path / "queries.txt"
    # BOM — так сохраняет Блокнот Windows
    path.write_text("﻿купить духи\n\n# комментарий\n  парфюм магазин  \nкупить духи\n",
                    encoding="utf-8")
    assert main.read_queries(["купить парфюм"], str(path)) == [
        "купить парфюм", "купить духи", "парфюм магазин"]


@pytest.mark.parametrize("encoding", ["cp1251", "utf-16"])
def test_read_queries_windows_encodings(tmp_path, encoding):
    # cp1251 — «Текст» из Excel в русской Windows, utf-16 — «Текст Юникод»
    path = tmp_path / "queries.txt"
    path.write_bytes("купить духи\r\nпарфюм магазин\r\n".encode(encoding))
    assert main.read_queries([], str(path)) == ["купить духи", "парфюм магазин"]


def test_read_queries_from_stdin(monkeypatch):
    stdin = io.TextIOWrapper(io.BytesIO("купить духи\nпарфюм магазин\n".encode("utf-8")))
    monkeypatch.setattr(sys, "stdin", stdin)
    assert main.read_queries([], "-") == ["купить духи", "парфюм магазин"]


# ---------- режим с аргументами ----------

def test_cli_json_paginates_until_ads_end(monkeypatch, tmp_path):
    session = FakeSession([ads_html("a.ru", "b.ru"), ads_html("c.ru", "d.ru")])
    out = tmp_path / "out.json"
    assert run_cli(monkeypatch, session, "купить духи", "-p", "5", "-r", "213",
                   "-o", str(out)) == main.EXIT_OK

    # третья страница — «ничего не нашлось», дальше не ходим
    assert [call["p"] for call in session.calls] == [0, 1, 2]
    assert session.calls[0] == {"text": "купить духи", "p": 0, "lr": 213}
    data = read_json(out)
    assert (data["region"], data["runs"]) == (213, 1)
    assert data["queries"] == [{"query": "купить духи", "status": "ok", "runs_with_ads": 1,
                                "pages": 3, "ads": 4, "message": ""}]
    assert [(ad["domain"], ad["page"], ad["position"], ad["seen"]) for ad in data["ads"]] == [
        ("a.ru", 1, 1, 1), ("b.ru", 1, 2, 1), ("c.ru", 2, 3, 1), ("d.ru", 2, 4, 1)]


def test_cli_region_by_name(monkeypatch, tmp_path):
    session = FakeSession([ADS_PAGE])
    run_cli(monkeypatch, session, "создание сайтов", "-r", "Ставрополь",
            "-o", str(tmp_path / "out.json"))
    assert session.calls[0]["lr"] == 36


def test_cli_json_to_stdout(monkeypatch, capsys):
    assert run_cli(monkeypatch, FakeSession([ADS_PAGE]), "купить духи", "-o", "-") == main.EXIT_OK
    data = json.loads(capsys.readouterr().out)
    assert len(data["ads"]) == 2
    assert data["ads"][0]["sitelinks"][0] == "Aroma Box"


def test_cli_xlsx_has_ads_and_queries_sheets(monkeypatch, tmp_path):
    out = tmp_path / "ads.xlsx"
    assert run_cli(monkeypatch, FakeSession([ADS_PAGE]), "купить духи", "-o", str(out)) == 0
    workbook = load_workbook(out)
    assert workbook.sheetnames == ["Объявления", "Запросы"]
    rows = list(workbook["Объявления"].values)
    assert rows[0] == tuple(main.AD_COLUMNS.values())
    assert rows[1][:6] == ("купить духи", 1, 1, 1, "randewoo.ru",
                           "Парфюм в интернет-магазине Randewoo — скидки до 55%")
    assert rows[1][7:9] == ("Aroma Box | Бестселлеры | Конструктор Aroma Box | SALE", "Промо")
    assert list(workbook["Запросы"].values)[1][:5] == ("купить духи", "ok", 1, 1, 2)


def test_cli_csv(monkeypatch, tmp_path):
    out = tmp_path / "ads.csv"
    assert run_cli(monkeypatch, FakeSession([ADS_PAGE]), "купить духи", "-o", str(out)) == 0
    with open(out, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    assert rows[0] == list(main.AD_COLUMNS.values())
    assert [row[4] for row in rows[1:]] == ["randewoo.ru", "impexoils.ru"]


def test_cli_default_output_goes_to_results(monkeypatch):
    assert run_cli(monkeypatch, FakeSession([ADS_PAGE]), "купить духи") == main.EXIT_OK
    assert len(list(main.RESULTS_DIR.glob("ads_*.xlsx"))) == 1


def test_cli_no_ads_exit_code(monkeypatch, tmp_path):
    out = tmp_path / "out.json"
    assert run_cli(monkeypatch, FakeSession([]), "редкий запрос", "-o", str(out)) == main.EXIT_NO_ADS
    query = read_json(out)["queries"][0]
    assert (query["status"], query["message"]) == ("no_ads", main.STATUS_MESSAGES["no_ads"])


def test_cli_captcha_exit_code(monkeypatch, tmp_path):
    session = FakeSession([ADS_PAGE], url="https://yandex.ru/showcaptcha?retpath=abc")
    out = tmp_path / "out.json"
    assert run_cli(monkeypatch, session, "купить духи", "-o", str(out)) == main.EXIT_ERRORS
    assert read_json(out)["queries"][0]["status"] == "captcha"


def test_cli_logs_request_estimate(monkeypatch, tmp_path, capsys):
    run_cli(monkeypatch, FakeSession([]), "а", "б", "-p", "3", "--runs", "2",
            "-o", str(tmp_path / "out.json"))
    assert "Запросов к Яндексу: до 12" in capsys.readouterr().err


@pytest.mark.parametrize("argv", [
    [], ["купить", "-o", "out.txt"], ["купить", "-p", "0"], ["купить", "-r", "Атлантида"],
    ["купить", "--runs", "0"], ["купить", "--runs", str(main.MAX_RUNS + 1)],
])
def test_cli_bad_arguments(argv):
    with pytest.raises(SystemExit) as exc:
        main.main(argv)
    assert exc.value.code == main.EXIT_USAGE


def test_cli_checks_output_before_requests(monkeypatch, tmp_path):
    # Так в Windows ведёт себя файл, открытый в Excel: писать в него нельзя
    locked = tmp_path / "locked.xlsx"
    locked.write_bytes(b"")
    locked.chmod(0o444)
    session = use_session(monkeypatch, FakeSession([ADS_PAGE]))
    for output in (str(locked), str(tmp_path / "нет-папки" / "out.xlsx")):
        with pytest.raises(SystemExit) as exc:
            main.main(["купить духи", "-o", output])
        assert exc.value.code == main.EXIT_USAGE
    assert session.calls == []


# ---------- прогоны ----------

def test_runs_are_merged(monkeypatch, tmp_path):
    session = QueueSession(ads_html("a.ru", "b.ru"), ads_html("b.ru", "c.ru"), NO_ADS_PAGE)
    out = tmp_path / "out.json"
    assert run_cli(monkeypatch, session, "окна", "--runs", "3", "-o", str(out)) == main.EXIT_OK

    data = read_json(out)
    assert data["runs"] == 3
    # b.ru встретилось дважды и во втором прогоне было выше — берётся лучшая позиция
    assert [(ad["domain"], ad["position"], ad["seen"]) for ad in data["ads"]] == [
        ("a.ru", 1, 1), ("b.ru", 1, 2), ("c.ru", 2, 1)]
    assert data["queries"][0] == {"query": "окна", "status": "ok", "runs_with_ads": 2,
                                  "pages": 1, "ads": 3, "message": ""}


def test_runs_go_through_all_queries_before_repeating(monkeypatch, tmp_path):
    session = FakeSession([ADS_PAGE])
    run_cli(monkeypatch, session, "первый", "второй", "--runs", "2",
            "-o", str(tmp_path / "out.json"))
    assert [call["text"] for call in session.calls] == ["первый", "второй", "первый", "второй"]


def test_same_ad_on_two_pages_counts_once_per_run(monkeypatch, tmp_path):
    session = FakeSession([ads_html("a.ru", "b.ru"), ads_html("a.ru", "c.ru")])
    out = tmp_path / "out.json"
    run_cli(monkeypatch, session, "окна", "-p", "2", "-o", str(out))
    assert [(ad["domain"], ad["position"], ad["seen"]) for ad in read_json(out)["ads"]] == [
        ("a.ru", 1, 1), ("b.ru", 2, 1), ("c.ru", 4, 1)]


def test_captcha_in_later_run_keeps_collected_ads(monkeypatch, tmp_path):
    session = QueueSession(ads_html("a.ru"), CAPTCHA_PAGE)
    out = tmp_path / "out.json"
    assert run_cli(monkeypatch, session, "окна", "--runs", "2", "-o", str(out)) == main.EXIT_ERRORS
    data = read_json(out)
    assert (data["queries"][0]["status"], data["queries"][0]["runs_with_ads"]) == ("captcha", 1)
    assert [ad["domain"] for ad in data["ads"]] == ["a.ru"]


# ---------- цвет ----------

def test_paint_colors_only_real_terminals(monkeypatch):
    monkeypatch.delenv("NO_COLOR", raising=False)
    monkeypatch.setenv("TERM", "xterm-256color")
    monkeypatch.setattr(main, "_ansi", True)
    assert main.paint("ок", "bold", "green", stream=FakeTTY()) == "\033[1;32mок\033[0m"
    assert main.paint("ок", "green", stream=io.StringIO()) == "ок"  # файл или агент
    monkeypatch.setenv("NO_COLOR", "1")
    assert main.paint("ок", "green", stream=FakeTTY()) == "ок"


def test_windows_console_without_color_support_falls_back_to_plain_text(monkeypatch):
    # На Linux ctypes.windll нет — так же, как если включить цвет в консоли не удалось
    monkeypatch.setattr(main.sys, "platform", "win32")
    monkeypatch.setattr(main, "_ansi", True)
    main._enable_ansi()
    assert main._ansi is False


# ---------- запуск отдельным процессом: кодировки и зависимости ----------

def run_script(*args, python_flags=()):
    # cp1252 — как в Windows, когда вывод программы перехватывает агент или файл
    env = {**os.environ, "PYTHONIOENCODING": "cp1252"}
    return subprocess.run([sys.executable, *python_flags, str(ROOT / "main.py"), *args],
                          stdin=subprocess.DEVNULL, capture_output=True, env=env, timeout=60)


def test_redirected_output_is_utf8_whatever_the_system_encoding():
    result = run_script()  # без запросов и без терминала: ошибка, а не пошаговый режим
    assert result.returncode == main.EXIT_USAGE
    stderr = result.stderr.decode("utf-8")
    assert "не задано ни одного запроса" in stderr
    assert "\033[" not in stderr  # и без цветовых кодов


def test_missing_dependency_is_explained():
    result = run_script("купить духи", python_flags=("-S",))  # -S: без site-packages
    assert result.returncode == 1
    assert "pip install -r requirements.txt" in result.stderr.decode("utf-8")


# ---------- пошаговый режим ----------

def test_interactive_typed_queries(monkeypatch, no_sleep):
    session = use_session(monkeypatch, FakeSession([ads_html("a.ru", "b.ru"),
                                                    ads_html("c.ru", "d.ru")]))
    ask = answers("купить духи", "парфюм магазин", "", "Ставрополь", "2", "", "итог")
    assert main.interactive(ask) == main.EXIT_OK

    assert [(c["text"], c["p"], c["lr"]) for c in session.calls] == [
        ("купить духи", 0, 36), ("купить духи", 1, 36),
        ("парфюм магазин", 0, 36), ("парфюм магазин", 1, 36)]
    # просто имя файла — в папку results, расширение .xlsx подставляется
    rows = list(load_workbook(main.RESULTS_DIR / "итог.xlsx")["Объявления"].values)
    assert len(rows) == 1 + 8


def test_interactive_reasks_and_uses_defaults(monkeypatch, capsys, no_sleep):
    session = use_session(monkeypatch, FakeSession([ADS_PAGE]))
    ask = answers("", "купить духи", "", "Атлантида", "", "0", "x", "", "11", "", "")
    assert main.interactive(ask) == main.EXIT_OK

    printed = capsys.readouterr().out
    assert "Нужен хотя бы один запрос" in printed
    assert "Не знаю региона «Атлантида»" in printed
    assert "Регион: Россия (225)" in printed
    assert "Нужно целое число от 1." in printed
    assert f"Нужно целое число от 1 до {main.MAX_RUNS}." in printed
    assert session.calls[0]["lr"] == 225 and len(session.calls) == 1
    assert len(list(main.RESULTS_DIR.glob("ads_*.xlsx"))) == 1


def test_interactive_runs(monkeypatch, no_sleep):
    session = use_session(monkeypatch, QueueSession(ads_html("a.ru"), ads_html("b.ru")))
    assert main.interactive(answers("окна", "", "", "", "2", "итог.json")) == main.EXIT_OK
    assert len(session.calls) == 2
    data = read_json(main.RESULTS_DIR / "итог.json")
    assert (data["runs"], [ad["domain"] for ad in data["ads"]]) == (2, ["a.ru", "b.ru"])


def test_interactive_queries_from_dragged_file(monkeypatch, tmp_path, no_sleep):
    session = use_session(monkeypatch, FakeSession([ADS_PAGE]))
    queries = tmp_path / "мои запросы.txt"
    queries.write_text("купить духи\nпарфюм магазин\n", encoding="utf-8")
    out = tmp_path / "out.json"
    # Windows при перетаскивании файла в окно берёт путь в кавычки
    ask = answers(f'"{queries}"', "213", "", "", str(out))
    assert main.interactive(ask) == main.EXIT_OK
    assert [c["text"] for c in session.calls] == ["купить духи", "парфюм магазин"]
    assert read_json(out)["region"] == 213


def test_clean_path_unescapes_macos_drag_and_drop():
    if sys.platform == "win32":
        pytest.skip("на Windows обратная косая черта — разделитель пути")
    assert main._clean_path(r"/Users/me/My\ Files/q.txt ") == "/Users/me/My Files/q.txt"


def test_interactive_eof_cancels(capsys):
    assert main.interactive(answers()) == main.EXIT_USAGE
    assert "Отменено" in capsys.readouterr().out
